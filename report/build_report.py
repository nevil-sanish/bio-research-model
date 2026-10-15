"""Builds report/Report_2025BCD0011_GeneExpressionClustering.docx from the Colab results in model/."""
import json
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
V1, V2, V3 = (ROOT / "model" / d for d in ("version1_baseline", "version2_cv_tuned", "version3_large_unseen_test"))
RES, FIG = V3 / "results", V3 / "figures"
m = json.loads((RES / "metrics.json").read_text())
m1, m2 = (json.loads((v / "results" / "metrics.json").read_text())["test"] for v in (V1, V2))
large = pd.read_csv(RES / "large_unseen_tests.csv")
te, tr = m["test"], m["train"]
cv = pd.read_csv(RES / "cv_tuning.csv")
cm = pd.read_csv(RES / "confusion_test.csv", index_col=0)
classes = [c.replace("pred ", "") for c in cm.columns]
pcs = te["best_n_pcs"]
pcs_txt = f"{pcs} principal components" if pcs else "no PCA (scaled genes used directly)"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.55); s.left_margin = s.right_margin = Inches(0.7)

def h(t, lvl=1): doc.add_heading(t, lvl)
def p(t): doc.add_paragraph(t)
def bullets(items):
    for i in items: doc.add_paragraph(i, style="List Bullet")
def table(df, fmt="{:.4f}"):
    t = doc.add_table(rows=1, cols=len(df.columns) + 1); t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text = ""
    for j, c in enumerate(df.columns): t.rows[0].cells[j + 1].text = str(c)
    for idx, row in df.iterrows():
        cells = t.add_row().cells; cells[0].text = str(idx)
        for j, v in enumerate(row): cells[j + 1].text = fmt.format(v) if isinstance(v, float) else str(v)
def fig(name, cap, w=6.2):
    doc.add_picture(str(FIG / name), width=Inches(w))
    c = doc.add_paragraph(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic = True

t = doc.add_heading("Gene Expression Clustering and Analysis using K-Means", 0)
p("ISC 211 – Introduction to Bioinformatics · Individual Assignment\n"
  "Name: Nevil Mathew Sanish   ·   Roll No.: 2025BCD0011   ·   Indian Institute of Information Technology Kottayam")

h("1. Assignment allocation")
p("R = last three digits of roll number = 011 = 11.  Topic = ((11 − 1) mod 11) + 1 = 11 → “Gene Expression Clustering "
  "and Analysis” (Georgieva, ERIS 2022, CEUR-WS Vol-3372 paper02). Because Topic > 10 the algorithm is chosen from "
  "{M11, M12, M13}; M11 – K-Means clustering – was selected, as it is the objective-function clustering method the paper "
  "reviews (§2.2) and is directly suited to discovering groups in expression data.")

h("2. Research paper and biological problem")
p("The paper reviews clustering for gene-expression analysis – hierarchical, objective-function (k-means, fuzzy c-means, SOM) "
  "and density-based (DBSCAN) methods – and extends an iterative DBSCAN procedure for detecting differentially expressed genes "
  "between two mouse strains (C57BL/6J vs DBA/2J, Bottomly et al. RNA-seq). Key points used in this work:")
bullets([
    "Biological problem: expression tables contain thousands of genes for many samples; unsupervised clustering is needed to "
    "discover unknown partitions – groups of samples/genes with similar behaviour, e.g. distinct cancer types (paper §1).",
    "Pre-processing: a log transformation is required because some genes are expressed orders of magnitude higher than others, "
    "and non-active (zero) genes must be filtered (paper §3.2).",
    "Choosing the number of clusters k is the main difficulty of k-means; PCA-based representations help (paper §2.2).",
])
p("Here the same problem – discovering biologically meaningful groups from RNA-seq expression without using labels – is "
  "addressed with K-Means on tumour samples, with the known tumour type used only to evaluate the discovered clusters.")

h("3. Dataset")
p("The paper’s mouse dataset has no ground-truth groups for clustering evaluation, so a suitable public dataset with known "
  "biology was selected: UCI “Gene Expression Cancer RNA-Seq” (UCI ML Repository, ID 401), a random extract of the TCGA "
  "Pan-Cancer Atlas (Illumina HiSeq). It was downloaded from "
  "https://archive.ics.uci.edu/static/public/401/gene+expression+cancer+rna+seq.zip (a zip holding "
  "TCGA-PANCAN-HiSeq-801x20531.tar.gz with data.csv and labels.csv); a copy is in the dataset/ folder.")
bullets([
    "801 tumour samples × 20,531 genes (RNA-seq expression values).",
    "5 tumour types: BRCA (breast, 300), KIRC (kidney, 146), LUAD (lung, 141), PRAD (prostate, 136), COAD (colon, 78).",
])

h("4. Data preparation")
bullets([
    "log2(x + 1) transform of all expression values (paper §3.2).",
    "Removal of zero-variance (non-active) genes → 20,264 genes.",
    "Stratified 80/20 split (seed 42): 640 training samples, 161 unseen test samples. Everything below is fitted on the "
    "training set only, so no information from the test set leaks into the model.",
    "Selection of the N most variable genes, per-gene standardisation (z-score) and optional PCA – N and the number of PCA "
    "components are hyper-parameters chosen by cross-validation.",
])

h("5. Algorithm and training – K-Means (M11)")
p("K-Means partitions samples into k clusters by iteratively assigning each sample to its nearest centroid (Euclidean distance) "
  "and recomputing centroids, minimising the within-cluster sum of squares (inertia). k = 5 was used, matching the five tumour "
  "types and supported by the elbow/silhouette analysis (Fig. 1). Each fit uses 50 random initialisations (k-means++) and keeps the "
  "lowest-inertia solution. Because cluster IDs are arbitrary, each cluster is mapped to a tumour type with the Hungarian "
  "algorithm using the training labels only; this mapping is then applied unchanged to unseen samples.")
p(f"Hyper-parameter tuning: 20 pipelines (top-variance genes ∈ {{500, 1000, 2000, 5000}} × PCA ∈ {{none, 5, 10, 20, 50}}) "
  f"were scored with 5-fold stratified cross-validation inside the training set. The best was {te['best_n_genes']} genes with "
  f"{pcs_txt}. Top configurations:")
top = cv.head(5).copy(); top.index = [f"#{i+1}" for i in range(len(top))]
top.columns = ["genes", "PCA comps", "CV accuracy (mean)", "CV accuracy (std)"]
table(top)
fig("k_selection.png", "Fig. 1 – Elbow (inertia) and silhouette versus k on the training set.", 4.3)

h("6. Testing on unseen data and evaluation")
p(f"The final model was trained once on all 640 training samples and evaluated once on the 161 held-out test samples. "
  f"Accuracy on unseen data = {te['accuracy']*100:.2f}% ({te['n_correct']}/{te['n_test']} samples correctly grouped). "
  "Metrics: accuracy (after training-learned cluster→class mapping), Adjusted Rand Index (ARI, 1 = perfect agreement, 0 = random), "
  "Normalised Mutual Information (NMI), and silhouette coefficient (cluster compactness/separation, −1…1).")
mt = pd.DataFrame({k: {"Accuracy": d["accuracy"], "ARI": d["ARI"], "NMI": d["NMI"], "Silhouette": d["silhouette"]}
                   for k, d in [("Training (seen)", tr), ("Test (unseen)", te)]}).T
table(mt)
doc.add_paragraph()
vals = cm.values; prec = vals.diagonal() / vals.sum(0).clip(1); rec = vals.diagonal() / vals.sum(1).clip(1)
f1 = 2 * prec * rec / (prec + rec).clip(1e-9)
pr = pd.DataFrame({"Precision": prec, "Recall": rec, "F1-score": f1, "Support": [str(int(v)) for v in vals.sum(1)]}, index=classes)
table(pr)
fig("confusion_test.png", "Fig. 2 – Confusion matrix on the unseen test set.", 3.0)
fig("clusters_pca.png", "Fig. 3 – Unseen test samples in the first two PCs: K-Means clusters (left) vs true tumour types (right).", 5.3)

h("7. Testing on a large unseen dataset")
p("To check that the result does not depend on one small test split, the final pipeline (hyper-parameters fixed) was retrained "
  "on only 50 % or only 20 % of all 801 samples and tested on the remaining 401 or 641 never-seen samples, repeated over "
  "10 random stratified splits each (results/large_unseen_tests.csv).")
lg = large.groupby(["train_samples", "unseen_test_samples"]).agg(
    acc_mean=("accuracy", "mean"), acc_std=("accuracy", "std"), acc_min=("accuracy", "min"), ari_mean=("ARI", "mean"))
lg.index = [f"train {a} / unseen {b}" for a, b in lg.index]
lg.columns = ["Accuracy (mean)", "Accuracy (std)", "Accuracy (min)", "ARI (mean)"]
table(lg)
big = lg.loc[f"train {large.train_samples.min()} / unseen {large.unseen_test_samples.max()}"]
p(f"Even when trained on only 160 samples, the model groups 641 unseen samples with {big['Accuracy (mean)']*100:.2f}% mean accuracy, "
  "confirming that the tumour-type structure found by K-Means is stable and generalises to new data.")

h("8. Model versions")
p("Three versions were trained in Google Colab; each is stored in its own folder under model/ with its notebook (with outputs), "
  "saved model (.joblib), results (.csv/.json) and figures.")
vt = pd.DataFrame({
    "Changes": ["2,000 top-variance genes, k = 5; cluster→class matching done on the test set",
                f"5-fold CV tuning on training set → {m2['best_n_genes']} genes; cluster→class map learned on training labels",
                "Version 2 + robustness test on 401 / 641 unseen samples (final)"],
    "Unseen-test accuracy": [f"{m1['accuracy']*100:.2f}%", f"{m2['accuracy']*100:.2f}%", f"{te['accuracy']*100:.2f}%"],
    "ARI": [f"{m1['ARI']:.3f}", f"{m2['ARI']:.3f}", f"{te['ARI']:.3f}"],
}, index=["v1 – version1_baseline", "v2 – version2_cv_tuned", "v3 – version3_large_unseen_test"])
table(vt)
p("Improvement from v1 to v2 comes from (i) using fewer, more informative genes – 500 instead of 2,000 – which removes noisy "
  "dimensions from the Euclidean distance, and (ii) a stricter, leak-free evaluation where the cluster→class mapping is learned only "
  "from training data.")

h("9. Analysis of results")
errs = [(classes[i], classes[j], int(vals[i, j])) for i in range(len(classes)) for j in range(len(classes)) if i != j and vals[i, j]]
err_txt = "; ".join(f"{n} {a} sample(s) grouped with {b}" for a, b, n in errs) or "no errors"
bullets([
    f"Without ever seeing labels during clustering, K-Means recovered the five tumour types almost perfectly: "
    f"{te['accuracy']*100:.2f}% accuracy and ARI {te['ARI']:.3f} on unseen samples, close to the training result "
    f"({tr['accuracy']*100:.2f}%), so the model generalises and does not over-fit.",
    f"Misclassifications on the test set: {err_txt}. Errors occur between tissues with related biology/expression programmes "
    "(e.g. adenocarcinomas of different organs), while tissue-specific cancers such as KIRC and PRAD form very compact clusters.",
    "Cross-validation showed that restricting to the 500 most variable genes gave the best accuracy (99.84% CV vs 99.22% with 5,000 genes), while PCA "
    "gave no further gain: low-variance genes add noise to Euclidean distances, confirming the paper’s point that the proximity measure and "
    "the data representation determine cluster quality.",
    f"The silhouette score is moderate ({te['silhouette']:.2f}) even though accuracy is high: in high-dimensional expression space "
    "the clusters are well separated in direction but not perfectly spherical – consistent with the paper’s observation that "
    "gene-expression clusters are often elongated rather than Euclidean.",
    "Choosing k: the elbow/silhouette curves support k ≈ 5, agreeing with the five biological tumour types, illustrating the "
    "paper’s remark that the number of clusters must be determined by an additional search procedure.",
])

h("10. Conclusion")
p("Following the assigned paper, RNA-seq expression data were log-transformed and filtered, and K-Means (M11) was used to "
  "discover tumour-type clusters without supervision. With hyper-parameters chosen by cross-validation on training data only, the "
  f"model reaches {te['accuracy']*100:.2f}% accuracy (ARI {te['ARI']:.3f}, NMI {te['NMI']:.3f}) on unseen samples, showing "
  "that clustering of gene expression reveals biologically meaningful groups. Limitations: k was fixed using prior knowledge of "
  "the five classes, and K-Means assumes spherical clusters; hierarchical, density-based (DBSCAN) or Mahalanobis-distance "
  "clustering, as discussed in the paper, are natural extensions.")

h("11. Submission folder structure")
bullets([
    "dataset/ – original UCI archive (gene_expression_cancer_rna_seq_UCI401.zip: data.csv 801×20,531 + labels.csv) and labels.csv.",
    "model/version1_baseline/, model/version2_cv_tuned/, model/version3_large_unseen_test/ – notebook with outputs, model/kmeans_pancan.joblib, "
    "results/ (metrics.json, metrics.csv, confusion_test.csv, cv_tuning.csv, large_unseen_tests.csv) and figures/.",
    "model/build_notebook.py and model/kmeans_gene_expression.ipynb – source that generates the final notebook (run in Google Colab).",
    "report/ – this report and build_report.py, which builds it from the saved results.",
])

h("References")
for r in ["O. Georgieva, “Clustering for Gene Expression Analysis,” ERIS 2022, CEUR-WS Vol-3372, pp. 17–24, 2022.",
          "S. Fiorini, “Gene expression cancer RNA-Seq,” UCI Machine Learning Repository, 2016, doi:10.24432/C5R88H.",
          "The Cancer Genome Atlas Research Network et al., “The Cancer Genome Atlas Pan-Cancer analysis project,” Nat. Genet. 45, 2013.",
          "F. Pedregosa et al., “Scikit-learn: Machine Learning in Python,” JMLR 12, 2011."]:
    doc.add_paragraph(r, style="List Number")

out = ROOT / "report" / "Report_2025BCD0011_GeneExpressionClustering.docx"
doc.save(out); print("wrote", out)
